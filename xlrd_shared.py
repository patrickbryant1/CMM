#! /usr/bin/env python
# -*- coding: utf-8 -*-

'''
This is a program that takes an xl workbook, reads it and uses its
sheets to perform filtering as specified by the text file.

-Remember that you cannot write too many rows to an excel sheet,
there is a limitation!
'''

import sys
import os
import pdb
import xlrd
import argparse
import openpyxl

#Arguments for argparse module:
parser = argparse.ArgumentParser(description = '''
This is a program that takes an xl workbook, reads it and uses its
sheets to perform filtering as specified by text files.

-Remember that you cannot write too many rows to an excel sheet,
there is a limitation!
''')
 
parser.add_argument('workbook_r', nargs=1, type= str,
                  default=sys.stdin, help = 'path to excel file with variants to be opened')


parser.add_argument('ref_dbs', nargs=1, type= str,
                  default=sys.stdin, help = '''path to text file containing column positions
                  for ref_dbs to filter on.''')

parser.add_argument('zygosity_positions', nargs=1, type= str,
                  default=sys.stdin, help = '''path to text file containing
                  zygosity_positions to filter on.''')

###########################################################################################
#Functions

def text_to_list(file_name):
    '''A function that reads a text file and creates a list of its rows.
    Input: text_file (txt)
    Output: filter_options (list)
    '''

    filter_options = [] #empty list to store options in
    with open(file_name, 'r') as infile:
        for line in infile:
            line = line.rstrip('\n')
            filter_options.append(line)
            
    return filter_options

def encode_ascii(xl_sheet, row_idx, col_idx):
    '''A function that encodes and retrieves excel cell values
    Input: row and column positions
    Output: cell_value
    '''
    cell_value = str(xl_sheet.cell(row_idx, col_idx).value).encode('ascii','ignore')
    return cell_value
    
def find_shared(zygosity_positions, xl_sheet, row_idx, above_t):
    '''Finds the variants shared between the samples as specified by zygosity_positions
        Input = zygosity_positions, xl_sheet, row_idx, above_t
        Output = shared (bool)
    '''
    
    share_pos = [] #List to store colummn positions for those that should share
    not_share_pos = [] #List to store colummn positions for those that should not share
    shared = True #See if the variants are shared

    #Get positions for zygosities
    for item in zygosity_positions:
        item = item.split(' ') #split on space
        if item[2] == '1':
            share_pos.append(item[1])  #get column position
        else:
            not_share_pos.append(item[1])

    wt = False #Keep track of wt
    hom = False #Keep track of hom
    other = False #Keeep track of other
    count = 0 #To keep track of how many oth
    for pos in share_pos:
        zyg = encode_ascii(xl_sheet, row_idx, int(pos)) #zygosity to match
        if zyg =='.':  #If the zygosity cannot be assessed, it is disregarded
            shared = False            
            break   
        else:
            if zyg == "wt" or not zyg: #If the zyg is empty, it is wt
                wt = True
            if zyg == "hom":
                hom = True
            if zyg == "het":
                continue
            if zyg == 'oth':
                other = True
                count+=1 #Keep track of how many oth
        
    if count != len(share_pos) and other == True: #If all that should share are not oth
        shared = False
    if wt == True and hom == True: #If both hom and wt is true, the variant is not shared
        shared = False
    if wt == True and above_t == False: #If the variant is wt, the MAF should be above the threshold
        shared = False
    if hom == True and above_t == True: #If the variant is hom, the MAF should be below the threshold
        shared = False

    if shared == True and not_share_pos:
        shared = not_shared(not_share_pos, wt, hom, other, shared, above_t, xl_sheet, row_idx)

    return(shared)
            
def not_shared(not_share_pos, wt, hom, other, shared, above_t, xl_sheet, row_idx):
    #Check the ones that should not share
    for pos in not_share_pos:
        zyg = encode_ascii(xl_sheet, row_idx, int(pos)) #zygosity to match
        if zyg =='.': #If the zygosity cannot be assessed, it is disregarded
            shared = False
            break
        else:
            if zyg == 'het': #If the zygosity is het, they will share
                shared = False
            if zyg == "wt" or not zyg: #If the zyg is empty it is wt
                if wt == True or above_t == True: #If the variant is wt, the MAF should be below the threshold
                    shared = False
            if zyg == "hom": #If the variant is hom, the MAF should be above the threshold
                if hom == True or above_t == False:
                    shared = False 
            if other == True and zyg =='oth':
                shared = False
                
    return(shared)

def filter_sheet(workbook_r, name, ref_dbs, zygosity_positions):        
    '''A function that takes an xl-sheet
    as input and finds the shared variants as specified in zygosity positions
    Input: workbook_r, name, ref_dbs, zygosity_positions
    Output: None
    '''
    
    #Create an excel workbook and a sheet to write to
    workbook_w = openpyxl.Workbook()
    sheet_w = workbook_w.active #Get active sheet
    sheet_w.title = 'maf_filtered'

    shared = False

    for num in range(0, 1):#workbook_r.nsheets):
        sheet_r = workbook_r.sheet_by_index(num) #Open sheet_num
	
        row_idx = 0 #Row to transfer from
        row_n = 1 #To keep track of which row to write to
        #transfer first row
        write_to_sheet(row_idx, sheet_w, sheet_r, row_n)


        for row_idx in range(1, sheet_r.nrows):
            above_t = False
 	    for ref_db in ref_dbs:
                ref_db = ref_db.split(' ') #split each list item on space
                MAF = encode_ascii(sheet_r, row_idx, int(ref_db[1]))#filter on reference databases MAFs
                if MAF and MAF != '.': #Checks if MAF is empty or NA
                    MAF = float(MAF)
                    if MAF > float(ref_db[3]): #Filtering on MAF above threshold
                        above_t = True
            
            shared = find_shared(zygosity_positions, sheet_r, row_idx, above_t)
            if shared == True:
                
                    row_n += 1
                    write_to_sheet(row_idx, sheet_w, sheet_r, row_n)
                    
            else:
                continue

    workbook_w.save('shared_'+name[0])

    return None

def write_to_sheet(row_idx, sheet_w, sheet_r, row_n):
    '''A function that writes data into a sheet in the excel workbook.
    Input: row_idx, sheet_w, xl_sheet
    Output: None
    '''

    #Iterate over all columns
    for col_idx in range(0, sheet_r.ncols):
        cell = sheet_r.cell(row_idx, col_idx).value
        if type(cell) == float:
            	cell = str(cell) 
        cell = cell.encode('ascii','ignore')
        sheet_w.cell(row = row_n, column = col_idx+1).value = cell #Openpyxl is not zero indexed, but xlrd is

    return None
    
    
#Main program
args = parser.parse_args()

try:
    workbook_r = xlrd.open_workbook(args.workbook_r[0])
    #The opened workbook is closed automatically after use
    
    #Checks if file is empty
    if os.stat(args.workbook_r[0]).st_size == 0:
       print "Empty file."       
except IOError:
    print 'Cannot open', args.workbook_r[0]
else:


    #Arguments to filter on
    ref_dbs = text_to_list(args.ref_dbs[0])
    zygosity_positions = text_to_list(args.zygosity_positions[0])
    
    #Perform the filtering
    name = str(args.workbook_r[0]).split('/')[-1:] #name is only the last part of path
                                            #name is a list here
    filter_sheet(workbook_r, name, ref_dbs, zygosity_positions)


    
